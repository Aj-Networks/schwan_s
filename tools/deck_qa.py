import sys
from pptx import Presentation
from pptx.util import Emu
EMU=914400
p=Presentation(sys.argv[1])
W,H=p.slide_width/EMU,p.slide_height/EMU
# rough advance widths as a fraction of font size
ADV={"Bahnschrift SemiBold":0.52,"Bahnschrift Light":0.47,"Calibri":0.47,"Consolas":0.55}
issues=0
for i,s in enumerate(p.slides,1):
    boxes=[]
    for sh in s.shapes:
        if sh.left is None: continue
        x,y,w,h=[v/EMU for v in (sh.left,sh.top,sh.width,sh.height)]
        if sh.has_text_frame and sh.text.strip():
            # bounds: text must sit on the canvas
            if x< -0.05 or y< -0.05 or x+w>W+0.05 or y+h>H+0.05:
                print(f"  s{i} OUT OF BOUNDS: '{sh.text[:34]}' x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f}"); issues+=1
            # crude fit test, longest line against box width and height
            for para in sh.text_frame.paragraphs:
                size=None; face=None; txt="".join(r.text for r in para.runs)
                for r in para.runs:
                    if r.font.size: size=r.font.size.pt
                    if r.font.name: face=r.font.name
                if not txt or not size: continue
                adv=ADV.get(face,0.5)
                est=len(txt)*size*adv/72.0
                lines=max(1,round(est/max(w,0.1)+0.49))
                need=lines*size*1.35/72.0
                if est>w*1.02 and lines==1:
                    print(f"  s{i} TOO WIDE: '{txt[:34]}' needs {est:.2f}in in {w:.2f}in"); issues+=1
                if need>h*1.15:
                    print(f"  s{i} TOO TALL: '{txt[:30]}' ~{lines} lines need {need:.2f}in in {h:.2f}in"); issues+=1
            boxes.append((x,y,w,h,sh.text[:26]))
    # overlap between text boxes
    for a in range(len(boxes)):
        for b in range(a+1,len(boxes)):
            ax,ay,aw,ah,at=boxes[a]; bx,by,bw,bh,bt=boxes[b]
            ox=min(ax+aw,bx+bw)-max(ax,bx); oy=min(ay+ah,by+bh)-max(ay,by)
            if ox>0.25 and oy>0.18:
                print(f"  s{i} OVERLAP: '{at}' and '{bt}' ({ox:.2f} x {oy:.2f} in)"); issues+=1
print("slides:",len(p.slides._sldIdLst),"| issues:",issues)
